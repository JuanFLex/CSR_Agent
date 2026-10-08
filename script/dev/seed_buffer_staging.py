"""Rebuild a local stand-in for CSR_BUFFER_READINESS from CS_Trend_Forecast.xlsx.

Does in Python what script/sql/csr_buffer_readiness.sql does on the server:
keeps the "Balance vs SS (Buffer Stock)" row of every part, turns its 51 week
columns into rows, and dates them from the first real date of the commit
sheet, seven days apart. The stand-in has the table's real column names, so
the ColumnMap and the ingest run for real.

    python3 script/dev/seed_buffer_staging.py
    CSR_REPORTING_SQLITE=storage/development_staging.sqlite3 \
      bin/rails runner 'pp Csr::Buffer::Ingest.new(force: true).call.message'
"""
import openpyxl, sqlite3, datetime, os

SRC = "/mnt/c/Users/gdlguzju/Desktop/CS_Trend_Forecast.xlsx"
DB  = "/home/p_infinex/code/CSR_Agent/storage/development_staging.sqlite3"

COLS = ["FCP_PART", "FCP_SITE", "CPN", "MANUFACTURER_PART", "SUPPLIER_NAME", "CUSTOMER_ID",
        "DELIVERY_ADDRESS", "SAFETY_STOCK_CONSIDER_SPQ", "BUCKET_DATE", "VS_SS", "LOADED_AT"]

wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

# The anchor: first real date of the commit sheet ('Past' and 'Future' are text).
commit = wb["Vertical Commit (Hide)"]
first_week = min(r[6] for r in commit.iter_rows(min_row=2, values_only=True)
                 if isinstance(r[6], datetime.datetime)).date()
assert first_week.weekday() == 5, f"{first_week} is not a Saturday"

sheet = wb.worksheets[0]  # 'CS - Buffer Readiness Projectio'
rows = sheet.iter_rows(min_row=3, values_only=True)
header = next(rows)
i = {name: header.index(name) for name in (
    "FCP Part", "FCP Site", "Manufacturer Part", "Supplier Name", "Customer ID", "Delivery Address",
    "CPN", "Safety Stock (Consider SPQ)", "Type", "Past")}
weeks = range(i["Past"] + 1, len(header))  # every column after Past is a week
# The header knows its own dates; they have to agree with the anchor.
assert header[weeks[0]] == first_week.strftime("%m-%d-%y"), (header[weeks[0]], first_week)

loaded_at = datetime.datetime.fromtimestamp(os.path.getmtime(SRC)).isoformat(sep=" ")
out = []
for r in rows:
    if r[i["FCP Part"]] is None or r[i["Type"]] != "Balance vs SS (Buffer Stock)":
        continue
    part = [r[i[c]] for c in ("FCP Part", "FCP Site", "CPN", "Manufacturer Part", "Supplier Name",
                              "Customer ID", "Delivery Address", "Safety Stock (Consider SPQ)")]
    for n, col in enumerate(weeks):
        bucket = (first_week + datetime.timedelta(weeks=n)).isoformat()
        out.append([None if v is None else str(v) for v in part] + [bucket, None if r[col] is None else str(r[col]), loaded_at])

con = sqlite3.connect(DB)
con.execute("DROP TABLE IF EXISTS CSR_BUFFER_READINESS")
con.execute("CREATE TABLE CSR_BUFFER_READINESS (%s)" % ", ".join(f'"{c}" TEXT' for c in COLS))
con.executemany("INSERT INTO CSR_BUFFER_READINESS VALUES (%s)" % ",".join("?" * len(COLS)), out)
con.execute("CREATE INDEX csr_buffer_readiness_bucket ON CSR_BUFFER_READINESS (BUCKET_DATE)")
con.commit()
print(f"CSR_BUFFER_READINESS simulada: {len(out)} filas, {len(out) // len(weeks)} partes x {len(weeks)} semanas")
print("primera semana:", first_week, "| loaded_at:", loaded_at)
con.close()
